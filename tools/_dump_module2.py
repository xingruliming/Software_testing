import sys, os, zipfile, re, io, json

OUT = r"E:\0_work\1shijian\Software_testing\test_results\_module2_dump.txt"
lines = []

def w(s=""):
    lines.append(str(s))

# ---------- xlsx: 用例清单 ----------
xlsx = r"E:\0_work\1shijian\Software_testing\deliverables\module2\VGGT模块二测试用例清单.xlsx"
w("=" * 70)
w("XLSX: VGGT模块二测试用例清单.xlsx")
w("=" * 70)
try:
    import openpyxl
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    for ws in wb.worksheets:
        w("[SHEET] %s  dims=%s  max_row=%s max_col=%s" % (ws.title, ws.dimensions, ws.max_row, ws.max_column))
        for r in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column):
            vals = []
            for c in r:
                v = c.value
                if v is None:
                    vals.append("")
                else:
                    vals.append(str(v).replace("\n", " / ").strip())
            if any(vals):
                w("  R%02d | %s" % (r[0].row, " || ".join(vals)))
except Exception as e:
    w("XLSX ERROR: %r" % (e,))

# ---------- docx: 测试报告 / 缺陷报告 ----------
def dump_docx(path, title):
    w("")
    w("=" * 70)
    w("DOCX: %s" % title)
    w("=" * 70)
    try:
        z = zipfile.ZipFile(path)
        xml = z.read("word/document.xml").decode("utf-8", "ignore")
        # 段落切分
        paras = re.findall(r"<w:p[ >].*?</w:p>", xml, re.S)
        for i, p in enumerate(paras):
            texts = re.findall(r"<w:t[^>]*>(.*?)</w:t>", p, re.S)
            txt = "".join(texts)
            txt = (txt.replace("&amp;", "&").replace("&lt;", "<").replace("&gt;", ">")
                      .replace("&quot;", '"').replace("&apos;", "'"))
            txt = txt.strip()
            if txt:
                w("P%03d | %s" % (i, txt[:600]))
    except Exception as e:
        w("DOCX ERROR: %r" % (e,))

dump_docx(r"E:\0_work\1shijian\Software_testing\deliverables\module2\VGGT模块二测试报告.docx", "VGGT模块二测试报告")

with open(OUT, "w", encoding="utf-8") as f:
    f.write("\n".join(lines))
print("WROTE", OUT, len(lines), "lines")
