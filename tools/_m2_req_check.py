# -*- coding: utf-8 -*-
"""复核模块二需求清单：内容完整性与行高估算是否充足。"""
import openpyxl

TARGET = r"E:\0_work\1shijian\Software_testing\deliverables\module2\软件需求清单_模块二.xlsx"
M1 = r"E:\0_work\1shijian\Software_testing\deliverables\module1\软件需求清单_模块一.xlsx"
OUT = r"E:\0_work\1shijian\Software_testing\test_results\_m2_req_check.txt"

lines = []


def load(path):
    return openpyxl.load_workbook(path)["软件需求清单"]


ws = load(TARGET)
print_rows = []
lines.append("=== 模块二需求清单 ===")
lines.append("dims=%s  max_row=%s max_col=%s" % (ws.dimensions, ws.max_row, ws.max_column))
widths = {c: ws.column_dimensions[c].width for c in "ABCD"}
lines.append("列宽: %s" % widths)

for r in range(1, ws.max_row + 1):
    vals = [ws.cell(row=r, column=c).value for c in range(1, 5)]
    h = ws.row_dimensions[r].height
    # 估算 C 列所需行数与实际可用行数
    c_text = str(vals[2] or "")
    d_text = str(vals[3] or "")
    units_c = sum(1.0 if ord(ch) > 0x2E80 else 0.55 for ch in c_text)
    units_d = sum(1.0 if ord(ch) > 0x2E80 else 0.55 for ch in d_text)
    need_c = int(units_c / max(1.0, (widths["C"] or 62.7) - 1.5)) + 1
    need_d = int(units_d / max(1.0, (widths["D"] or 43.7) - 1.5)) + 1
    need = max(need_c, need_d, 1)
    fits = (h or 0) >= need * 17.5
    lines.append(
        "R%02d h=%-6s 需 %d 行(C) / %d 行(D) → %s | %s | %s | %s"
        % (r, h, need_c, need_d, "OK" if fits else "偏紧",
           str(vals[0] or "")[:22], str(vals[1] or "")[:24], (c_text or "")[:30])
    )
    print_rows.append((r, h, need, fits))

bad = [p for p in print_rows if not p[3]]
lines.append("")
lines.append("偏紧行: %s" % ([p[0] for p in bad] or "无"))

# 对比模块一行高
ws1 = load(M1)
lines.append("")
lines.append("=== 模块一行高参考 ===")
for r in range(2, ws1.max_row + 1):
    lines.append("R%02d h=%s" % (r, ws1.row_dimensions[r].height))

open(OUT, "w", encoding="utf-8").write("\n".join(lines))
print("WROTE", OUT)
