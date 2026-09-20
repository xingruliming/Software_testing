# -*- coding: utf-8 -*-
"""导出附录1模板与模块一需求清单的结构，供模块二需求清单对齐口径。"""
import openpyxl

TEMPLATE = r"E:\0_work\1shijian\实践作业-文档模板2025\附录1：软件需求清单模板.xlsx"
MODULE1 = r"E:\0_work\1shijian\Software_testing\deliverables\module1\软件需求清单_模块一.xlsx"
OUT = r"E:\0_work\1shijian\Software_testing\test_results\_m2_req_dump.txt"

lines = []


def dump(path, title):
    lines.append("#" * 78)
    lines.append("# %s" % title)
    lines.append("# %s" % path)
    lines.append("#" * 78)
    try:
        wb = openpyxl.load_workbook(path)
    except Exception as exc:
        lines.append("打开失败: %r" % (exc,))
        lines.append("")
        return
    for ws in wb.worksheets:
        lines.append("")
        lines.append("[SHEET] %s  dims=%s  max_row=%s max_col=%s" % (ws.title, ws.dimensions, ws.max_row, ws.max_column))
        lines.append("  merged: %s" % (", ".join(str(r) for r in ws.merged_cells.ranges)[:600] or "(无)"))
        widths = {}
        for key, dim in (ws.column_dimensions or {}).items():
            if dim.width:
                widths[key] = round(dim.width, 1)
        lines.append("  col widths: %s" % widths)
        for r in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 60), max_col=ws.max_column):
            vals = []
            for c in r:
                v = c.value
                vals.append("" if v is None else str(v).replace("\n", " / ").strip())
            if any(vals):
                lines.append("  R%02d | %s" % (r[0].row, " | ".join(vals)))
    lines.append("")


dump(TEMPLATE, "附录1 模板")
dump(MODULE1, "模块一需求清单（交付件）")

open(OUT, "w", encoding="utf-8").write("\n".join(lines))
print("WROTE", OUT, len(lines), "lines")
