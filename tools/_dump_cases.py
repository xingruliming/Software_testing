# -*- coding: utf-8 -*-
"""Dump the module-1 test case workbook to text for inspection."""
import sys
import openpyxl

path = sys.argv[1]
wb = openpyxl.load_workbook(path, data_only=True)
print("SHEETS:", wb.sheetnames)
for ws in wb.worksheets:
    print("=" * 90)
    print("SHEET:", ws.title, "dims:", ws.dimensions, ws.max_row, "x", ws.max_column)
    for r in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 60), values_only=True):
        cells = ["" if c is None else str(c).replace("\n", "\\n") for c in r]
        line = " | ".join(cells).rstrip(" |")
        if line.strip():
            print(line[:600])
